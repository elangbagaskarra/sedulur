import csv

from datetime import datetime
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import render
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Font

from apps.invoices.models import Invoice
from apps.payments.models import Payment


ISSUED_STATUSES = [
    "UNPAID",
    "PARTIAL",
    "PAID",
    "OVERDUE",
]


def parse_date(value, default):
    try:
        return datetime.strptime(
            value,
            "%Y-%m-%d",
        ).date()

    except (TypeError, ValueError):
        return default


def get_report_dates(request):
    today = timezone.localdate()

    default_start = today.replace(
        day=1
    )

    start_date = parse_date(
        request.GET.get("start"),
        default_start,
    )

    end_date = parse_date(
        request.GET.get("end"),
        today,
    )

    return start_date, end_date


def get_report_data(start_date, end_date):

    invoices = (
        Invoice.objects
        .filter(
            invoice_date__range=[
                start_date,
                end_date,
            ],
            status__in=ISSUED_STATUSES,
        )
        .select_related("customer")
        .order_by(
            "-invoice_date",
            "-id",
        )
    )

    payments = (
        Payment.objects
        .filter(
            payment_date__range=[
                start_date,
                end_date,
            ]
        )
        .select_related(
            "invoice",
            "invoice__customer",
        )
        .order_by(
            "-payment_date",
            "-id",
        )
    )

    total_sales = (
        invoices.aggregate(
            total=Sum("grand_total")
        )["total"]
        or Decimal("0")
    )

    total_tax = (
        invoices.aggregate(
            total=Sum("tax_total")
        )["total"]
        or Decimal("0")
    )

    total_withholding = (
        invoices.aggregate(
            total=Sum(
                "withholding_total"
            )
        )["total"]
        or Decimal("0")
    )

    total_payments = (
        payments.aggregate(
            total=Sum("amount")
        )["total"]
        or Decimal("0")
    )

    receivables = (
        Invoice.objects
        .filter(
            balance_due__gt=0,
            status__in=[
                "UNPAID",
                "PARTIAL",
                "OVERDUE",
            ],
        )
        .aggregate(
            total=Sum("balance_due")
        )["total"]
        or Decimal("0")
    )

    return {
        "invoices": invoices,
        "payments": payments,
        "total_sales": total_sales,
        "total_payments": total_payments,
        "receivables": receivables,
        "total_tax": total_tax,
        "total_withholding":
            total_withholding,
    }


@login_required
def report_index(request):

    start_date, end_date = (
        get_report_dates(request)
    )

    data = get_report_data(
        start_date,
        end_date,
    )

    context = {
        "start_date": start_date,
        "end_date": end_date,
        **data,
    }

    return render(
        request,
        "reports/index.html",
        context,
    )


@login_required
def export_csv(request):

    start_date, end_date = (
        get_report_dates(request)
    )

    data = get_report_data(
        start_date,
        end_date,
    )

    response = HttpResponse(
        content_type="text/csv",
    )

    filename = (
        f"laporan-penjualan-"
        f"{start_date}-"
        f"{end_date}.csv"
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; '
        f'filename="{filename}"'
    )

    writer = csv.writer(response)

    writer.writerow([
        "Nomor Invoice",
        "Tanggal",
        "Pelanggan",
        "Status",
        "Subtotal",
        "Pajak",
        "PPh",
        "Grand Total",
        "Dibayar",
        "Sisa",
    ])

    for invoice in data["invoices"]:

        writer.writerow([
            invoice.invoice_number,
            invoice.invoice_date,
            invoice.customer_name,
            invoice.get_status_display(),
            invoice.subtotal,
            invoice.tax_total,
            invoice.withholding_total,
            invoice.grand_total,
            invoice.paid_amount,
            invoice.balance_due,
        ])

    return response


@login_required
def export_excel(request):

    start_date, end_date = (
        get_report_dates(request)
    )

    data = get_report_data(
        start_date,
        end_date,
    )

    workbook = Workbook()

    # ==========================
    # SHEET PENJUALAN
    # ==========================

    sales_sheet = workbook.active

    sales_sheet.title = (
        "Penjualan"
    )

    headers = [
        "Invoice",
        "Tanggal",
        "Pelanggan",
        "Status",
        "Subtotal",
        "Diskon",
        "PPN",
        "PPh",
        "Grand Total",
        "Dibayar",
        "Sisa",
    ]

    sales_sheet.append(headers)

    for cell in sales_sheet[1]:
        cell.font = Font(
            bold=True
        )

    for invoice in data["invoices"]:

        sales_sheet.append([
            invoice.invoice_number,
            invoice.invoice_date,
            invoice.customer_name,
            invoice.get_status_display(),
            float(invoice.subtotal),
            float(
                invoice.discount_total
            ),
            float(invoice.tax_total),
            float(
                invoice.withholding_total
            ),
            float(invoice.grand_total),
            float(invoice.paid_amount),
            float(invoice.balance_due),
        ])

    # ==========================
    # SHEET PEMBAYARAN
    # ==========================

    payment_sheet = (
        workbook.create_sheet(
            "Pembayaran"
        )
    )

    payment_headers = [
        "Tanggal",
        "Invoice",
        "Pelanggan",
        "Metode",
        "Referensi",
        "Nominal",
    ]

    payment_sheet.append(
        payment_headers
    )

    for cell in payment_sheet[1]:
        cell.font = Font(
            bold=True
        )

    for payment in data["payments"]:

        payment_sheet.append([
            payment.payment_date,
            payment.invoice.invoice_number,
            payment.invoice.customer_name,
            payment.get_method_display(),
            payment.reference_number,
            float(payment.amount),
        ])

    # ==========================
    # SHEET RINGKASAN
    # ==========================

    summary_sheet = (
        workbook.create_sheet(
            "Ringkasan"
        )
    )

    summary_sheet.append([
        "LAPORAN KEUANGAN"
    ])

    summary_sheet["A1"].font = Font(
        bold=True,
        size=14,
    )

    summary_sheet.append([
        "Periode",
        f"{start_date} s/d {end_date}",
    ])

    summary_sheet.append([])

    summary_sheet.append([
        "Total Penjualan",
        float(
            data["total_sales"]
        ),
    ])

    summary_sheet.append([
        "Total Pembayaran",
        float(
            data["total_payments"]
        ),
    ])

    summary_sheet.append([
        "Piutang Aktif",
        float(
            data["receivables"]
        ),
    ])

    summary_sheet.append([
        "PPN",
        float(
            data["total_tax"]
        ),
    ])

    summary_sheet.append([
        "PPh / Potongan",
        float(
            data[
                "total_withholding"
            ]
        ),
    ])

    # ==========================
    # WIDTH
    # ==========================

    for sheet in workbook.worksheets:

        for column in sheet.columns:

            max_length = 0

            column_letter = (
                column[0].column_letter
            )

            for cell in column:

                value = cell.value

                if value is None:
                    continue

                max_length = max(
                    max_length,
                    len(str(value)),
                )

            sheet.column_dimensions[
                column_letter
            ].width = min(
                max_length + 3,
                35,
            )

    # ==========================
    # RESPONSE
    # ==========================

    response = HttpResponse(
        content_type=(
            "application/"
            "vnd.openxmlformats-"
            "officedocument."
            "spreadsheetml.sheet"
        )
    )

    filename = (
        f"laporan-"
        f"{start_date}-"
        f"{end_date}.xlsx"
    )

    response[
        "Content-Disposition"
    ] = (
        f'attachment; '
        f'filename="{filename}"'
    )

    workbook.save(response)

    return response